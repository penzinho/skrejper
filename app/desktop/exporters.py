"""Write export rows to disk.

CSV keeps the exact format the existing runners produce — ``utf-8-sig`` (Excel
opens it without mangling ``č``/``ž``) and ``QUOTE_ALL`` — so files from the
desktop app and from ``scripts/run_*.py`` stay interchangeable.

XLSX is the desktop addition. ``openpyxl`` is imported lazily so that the worker
can still write CSV on a machine where the optional dependency is missing.
"""

import csv
from pathlib import Path


def write_csv(path: Path, rows: list[dict], fields: list[str]) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, quoting=csv.QUOTE_ALL)
        writer.writeheader()
        writer.writerows(rows)
    return path


def _column_width(field: str, rows: list[dict], limit: int = 60) -> int:
    widest = len(field)
    for row in rows:
        widest = max(widest, len(str(row.get(field) or "")))
        if widest >= limit:
            return limit
    return min(widest + 2, limit)


def write_xlsx(path: Path, rows: list[dict], fields: list[str], sheet_title: str = "Leads") -> Path:
    """Same columns as the CSV, with a frozen header row and readable widths."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    # Excel rejects sheet names over 31 chars or containing []:*?/\
    sheet.title = "".join(c for c in (sheet_title or "Leads") if c not in "[]:*?/\\")[:31] or "Leads"

    sheet.append(fields)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append([str(row.get(field) or "") for field in fields])

    sheet.freeze_panes = "A2"
    if rows:
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(fields))}{len(rows) + 1}"
    for index, field in enumerate(fields, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = _column_width(field, rows)

    workbook.save(path)
    return path


def xlsx_available() -> bool:
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        return False
    return True
