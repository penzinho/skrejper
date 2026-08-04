import sys
import tempfile
import unittest
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from app.desktop.exporters import write_csv, write_xlsx, xlsx_available
from app.desktop.pipeline import ARBEITSAGENTUR_FIELDS

ROWS = [
    {
        "email": "info@muster.de",
        "company": "Muster GmbH",
        "city": "München",
        "title": "Schweisser",
        "category": "Metall",
        "published_at": "2026-08-01",
        "detail_url": "https://example.invalid/1",
        "employer_website": "https://muster.de",
        "refnr": "REF-1",
        "source": "arbeitsagentur",
    }
]


class CsvTests(unittest.TestCase):
    def setUp(self):
        self._dir = tempfile.TemporaryDirectory()
        self.path = Path(self._dir.name) / "out.csv"
        self.addCleanup(self._dir.cleanup)

    def test_format_matches_the_existing_runners(self):
        # utf-8-sig + QUOTE_ALL, byte for byte, so files from the desktop app and
        # from scripts/run_*.py stay interchangeable.
        write_csv(self.path, ROWS, ARBEITSAGENTUR_FIELDS)
        raw = self.path.read_bytes()

        self.assertTrue(raw.startswith(b"\xef\xbb\xbf"), "missing UTF-8 BOM")
        text = raw.decode("utf-8-sig")
        header, first, _ = text.split("\r\n", 2)
        self.assertEqual(
            header,
            '"email","company","city","title","category",'
            '"published_at","detail_url","employer_website","refnr","source"',
        )
        self.assertTrue(first.startswith('"info@muster.de","Muster GmbH","München"'))

    def test_creates_missing_directories(self):
        nested = Path(self._dir.name) / "a" / "b" / "out.csv"
        write_csv(nested, ROWS, ARBEITSAGENTUR_FIELDS)
        self.assertTrue(nested.exists())

    def test_missing_keys_become_empty_cells(self):
        write_csv(self.path, [{"email": "a@b.hr"}], ARBEITSAGENTUR_FIELDS)
        text = self.path.read_text(encoding="utf-8-sig")
        self.assertIn('"a@b.hr","","",', text)


@unittest.skipUnless(xlsx_available(), "openpyxl not installed")
class XlsxTests(unittest.TestCase):
    def setUp(self):
        self._dir = tempfile.TemporaryDirectory()
        self.path = Path(self._dir.name) / "out.xlsx"
        self.addCleanup(self._dir.cleanup)

    def test_writes_the_same_columns_as_the_csv(self):
        from openpyxl import load_workbook

        write_xlsx(self.path, ROWS, ARBEITSAGENTUR_FIELDS, sheet_title="Metall")
        sheet = load_workbook(self.path).active

        self.assertEqual([c.value for c in sheet[1]], ARBEITSAGENTUR_FIELDS)
        self.assertEqual(sheet.cell(row=2, column=2).value, "Muster GmbH")
        self.assertEqual(sheet.cell(row=2, column=3).value, "München")
        self.assertEqual(sheet.freeze_panes, "A2")
        self.assertEqual(sheet.title, "Metall")

    def test_sheet_title_is_sanitised(self):
        from openpyxl import load_workbook

        write_xlsx(self.path, ROWS, ARBEITSAGENTUR_FIELDS, sheet_title="Bau/Ausbau: alles" * 4)
        title = load_workbook(self.path).active.title

        self.assertLessEqual(len(title), 31)
        self.assertNotIn("/", title)
        self.assertNotIn(":", title)

    def test_empty_result_still_produces_a_readable_file(self):
        from openpyxl import load_workbook

        write_xlsx(self.path, [], ARBEITSAGENTUR_FIELDS)
        sheet = load_workbook(self.path).active

        self.assertEqual([c.value for c in sheet[1]], ARBEITSAGENTUR_FIELDS)
        self.assertEqual(sheet.max_row, 1)


if __name__ == "__main__":
    unittest.main()
